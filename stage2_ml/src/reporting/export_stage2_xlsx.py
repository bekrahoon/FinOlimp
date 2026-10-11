"""Create a Russian-language, presentation-only Stage 2 model report."""
from __future__ import annotations

import json
import os
import re
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPORT_NAME = "finolimp_stage2_model_report.xlsx"
_BAD_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(name="Aptos", bold=True, color="FFFFFF")
_HEADER_ALIGNMENT = Alignment(vertical="center", wrap_text=True)
_BODY_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
_SOURCE_CLASSES = ("1С", "агент", "бот", "выяснить")


def _classifier_label(name: str | None) -> str:
    labels = {
        "HistGradientBoostingClassifier": "Градиентный бустинг (HistGradientBoostingClassifier)",
        "DummyClassifier": "Резервная константная модель (DummyClassifier)",
    }
    return labels.get(name, name or "не указана")


def _safe_value(value: Any) -> Any:
    if isinstance(value, str):
        return _BAD_XML.sub("\uFFFD", value)
    return value


def _write_sheet(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: list[list[Any]],
    *,
    percent_columns: set[int] | None = None,
) -> None:
    sheet = workbook.create_sheet(title)
    sheet.freeze_panes = "A2"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGNMENT
    for row in rows:
        sheet.append([_safe_value(value) for value in row])
        for cell in sheet[sheet.max_row]:
            cell.alignment = _BODY_ALIGNMENT
            # Report values can originate from raw input text; never treat them as formulas.
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.row_dimensions[1].height = 34
    sheet.auto_filter.ref = sheet.dimensions
    percent_columns = percent_columns or set()
    for column_index, column_cells in enumerate(sheet.columns, start=1):
        values = [str(cell.value) for cell in column_cells if cell.value is not None]
        width = min(max(max((len(value) for value in values), default=12) + 2, 14), 58)
        sheet.column_dimensions[get_column_letter(column_index)].width = width
        if column_index in percent_columns:
            for cell in column_cells[1:]:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.00%"


def _validation_rows(inference: dict[str, Any]) -> list[list[Any]]:
    status = inference.get("status_model", {}).get("validation_vs_silver", {})
    ticket = inference.get("ticket_model", {}).get("validation_at_threshold", {})
    payment = inference.get("payment_model", {}).get("validation_at_threshold", {})
    source_model = inference.get("source_model", {})
    source = source_model.get("validation_vs_silver", {})

    def row(name, metrics, precision_key, recall_key, note):
        return [
            name,
            metrics.get("n"),
            metrics.get("threshold"),
            metrics.get("accuracy"),
            metrics.get(precision_key),
            metrics.get(recall_key),
            metrics.get("macro_f1"),
            metrics.get("false_positives"),
            metrics.get("false_negatives"),
            note,
        ]

    rows = [
        row("Статус строк 1С", status, "keep_precision", "keep_recall",
            "Сопоставление с метками допуска Stage 1 («silver»)."),
        row("Пары билетов", ticket, "precision_vs_silver_target",
            "recall_vs_silver_target", "Метрики относительно меток Stage 1 («silver»)."),
        row("Пары платежей", payment, "precision_vs_silver_target",
            "recall_vs_silver_target", "Метрики относительно меток Stage 1 («silver»)."),
    ]
    if source_model.get("skipped_for_bootstrap"):
        rows.append(["Источник расхождения", None, None, None, None, None, None, None, None,
                     "Модель источника не обучалась в этом запуске."])
    elif source:
        rows.append([
            "Источник расхождения", source.get("n"), None, source.get("accuracy"),
            None, None, source.get("macro_f1"), None, None,
            "Точность и полнота приведены отдельно по классам на листе «По классам источника».",
        ])
    else:
        rows.append(["Источник расхождения", None, None, None, None, None, None, None, None,
                     "Метрики проверки источника отсутствуют."])
    return rows


def _candidate_gate_rows(inference: dict[str, Any]) -> list[list[Any]]:
    gate = inference.get("ticket_model", {}).get("candidate_gate_silver_validation", {})
    note = "Доля положительных пар Stage 1, прошедших предварительный фильтр; это не полнота итогового сопоставления."
    rows = []
    if gate.get("ticket_positive_pairs") is not None:
        rows.append([
            "Билеты: общий хвост номера",
            gate.get("ticket_positive_pairs"),
            gate.get("ticket_exact_tail_gate_covered"),
            gate.get("ticket_exact_tail_gate_recall_silver"),
            note,
        ])
    if gate.get("payment_positive_pairs") is not None:
        rows.append([
            "Платежи: сумма и дата",
            gate.get("payment_positive_pairs"),
            gate.get("payment_amount_date_gate_covered"),
            gate.get("payment_amount_date_gate_recall_silver"),
            note,
        ])
    if not rows:
        rows.append(["Нет данных", None, None, None, "Покрытие кандидатов не рассчитано."])
    return rows


def _source_class_rows(inference: dict[str, Any]) -> list[list[Any]]:
    model = inference.get("source_model", {})
    per_class = model.get("validation_vs_silver", {}).get("per_class", {})
    if model.get("skipped_for_bootstrap"):
        return [["Не обучалась", None, None, None, None]]
    rows = []
    classes = model.get("classes", list(_SOURCE_CLASSES))
    for name in classes:
        metrics = per_class.get(name, {})
        rows.append([
            name,
            metrics.get("support"),
            metrics.get("precision"),
            metrics.get("recall"),
            metrics.get("f1"),
        ])
    return rows or [["Нет метрик проверки", None, None, None, None]]


def _source_confusion(inference: dict[str, Any]) -> tuple[list[str], list[list[Any]]]:
    model = inference.get("source_model", {})
    classes = model.get("classes") or list(_SOURCE_CLASSES)
    matrix = model.get("validation_vs_silver", {}).get("confusion", [])
    headers = ["Фактический класс / прогноз модели", *classes]
    if len(matrix) != len(classes) or any(len(row) != len(classes) for row in matrix):
        label = "Модель не обучалась" if model.get("skipped_for_bootstrap") else "Нет данных"
        return headers, [[label, *([None] * len(classes))]]
    return headers, [[actual, *counts] for actual, counts in zip(classes, matrix)]


def _summary_rows(inference: dict[str, Any], training: dict[str, Any],
                  run_id: str) -> list[list[Any]]:
    split = inference.get("temporal_split", {})
    status_rows = split.get("status_rows", {})
    pair_rows = split.get("pair_rows", {})
    source = inference.get("source_model", {})
    status_model = inference.get("status_model", {})
    ticket_model = inference.get("ticket_model", {})
    payment_model = inference.get("payment_model", {})
    output = inference.get("output", {})
    fraction = split.get("validation_fraction_requested")
    months = split.get("validation_months_requested")
    if isinstance(fraction, (int, float)):
        share: Any = f"{fraction:.0%}"
    elif months:
        share = f"последние {int(months)} мес."
    else:
        share = fraction
    cutoff = split.get("cutoff")
    training_status = training.get("status", {})
    training_pair_split = training.get("pairs", {}).get("temporal_split", {})
    train_pair_rows = training_pair_split.get("train", {}).get("rows")
    validation_pair_rows = training_pair_split.get("validation", {}).get("rows")
    excluded_pair_rows = training_pair_split.get(
        "cross_period_or_unknown_date_excluded")
    return [
        ["Идентификатор запуска", run_id, "Результаты этого запуска; Stage 1 остаётся бухгалтерским эталоном."],
        ["Граница дат", cutoff, "Дата разделения обучающей и проверочной выборок."],
        ["Период обучения", f"до {cutoff}" if cutoff else split.get("train_dates"),
         "Операции до границы дат."],
        ["Период проверки", f"с {cutoff}" if cutoff else split.get("validation_dates"),
         "Поздние даты; не участвуют в обучении моделей."],
        ["Доля проверочной выборки", share,
         "Доля уникальных дат (по умолчанию 20%) или последние полные месяцы."],
        ["Уникальных дат", split.get("unique_dates"), "Даты текущих входов 1С и ETM."],
        ["Модель статуса строк", _classifier_label(status_model.get("classifier")),
         "Алгоритм определения включения строки в ML-сверку."],
        ["Модель пар билетов", _classifier_label(ticket_model.get("classifier")),
         "Алгоритм оценки пар строк с билетами."],
        ["Модель пар платежей", _classifier_label(payment_model.get("classifier")),
         "Алгоритм оценки пар платежей."],
        ["Модель источника расхождения",
         "пропущена" if source.get("skipped_for_bootstrap")
         else _classifier_label(source.get("classifier")),
         "Алгоритм классификации источника расхождения."],
        ["Строк статуса: обучение", status_rows.get("train", training_status.get("n_train")),
         "Примеры для обучения модели статуса."],
        ["Строк статуса: проверка", status_rows.get(
            "validation", training_status.get("n_validation")),
         "Примеры для проверки модели статуса."],
        ["Пар: обучение", pair_rows.get("train", train_pair_rows),
         "Пары для обучения моделей сопоставления."],
        ["Пар: проверка", pair_rows.get("validation", validation_pair_rows),
         "Пары для временной проверки моделей."],
        ["Пары исключены на границе", pair_rows.get(
            "cross_period_or_unknown_date_excluded", excluded_pair_rows),
          "Пересекают границу дат или имеют неизвестную дату."],
        ["Источник: строки обучения", source.get("training_rows"),
         "Строки с метками Stage 1 до границы дат."],
        ["Источник: строки проверки", source.get("validation_rows"),
         "Строки с метками Stage 1 после границы дат."],
        ["Расхождения в результате Stage 2", output.get("discrepancy_rows"),
         "Количество строк в CSV-отчёте Stage 2."],
        ["Метки", "Stage 1 («silver»)",
         "Производные от правил Stage 1; не независимая оценка человеком."],
    ]


def _notes_rows(inference: dict[str, Any]) -> list[list[Any]]:
    source = inference.get("source_model", {})
    notes = [
        ["Назначение", "Stage 2 — ML-второе мнение; оно не заменяет поручения и выводы Stage 1."],
        ["Метки", "Обучающая и проверочная выборки используют метки Stage 1 («silver»). Метрики показывают соответствие правилам, а не независимую точность совпадений."],
        ["Временная проверка", "Проверочная выборка состоит из более поздних дат и не участвует в обучении. Пары, пересекающие границу дат, исключены."],
        ["Статусные признаки", "Групповые агрегаты используют только более ранние даты."],
        ["Покрытие кандидатов", "Доля показывает, какая часть положительных пар Stage 1 попала в набор кандидатов; это не фактическая полнота поиска совпадений."],
        ["Матрица ошибок", "Строки — фактические метки Stage 1, столбцы — классы, предсказанные моделью."],
    ]
    if not source.get("skipped_for_bootstrap"):
        notes.append([
            "Уверенность модели источника",
            "Оценки модели не калиброваны; это баллы модели, а не вероятность правильной классификации.",
        ])
        per_class = source.get("validation_vs_silver", {}).get("per_class", {})
        rare = [f"{name}: {metrics.get('support', 0)}" for name, metrics in per_class.items()
                if metrics.get("support", 0) < 10]
        if rare:
            notes.append(["Мало примеров в проверке", "; ".join(rare)
                          + ". Метрики по этим классам нестабильны."])
    return notes


def build_stage2_workbook(inference: dict[str, Any], training: dict[str, Any],
                          run_id: str, output: Path) -> Path:
    """Build a Russian report from one run's inference and training manifests."""
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.is_symlink():
        raise ValueError("Файл Stage 2 Excel-отчёта не может быть символической ссылкой.")

    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(workbook, "Сводка",
                 ["Показатель", "Значение", "Пояснение"],
                 _summary_rows(inference, training, run_id))
    _write_sheet(
        workbook, "Валидация",
        ["Задача", "Примеров", "Порог", "Доля верных ответов", "Точность положительных прогнозов",
         "Полнота", "Макро-F1", "Ложноположительных", "Ложноотрицательных",
         "Пояснение"],
        _validation_rows(inference), percent_columns={4, 5, 6, 7},
    )
    _write_sheet(
        workbook, "Покрытие кандидатов",
        ["Фильтр кандидатов", "Положительных пар Stage 1", "Попало в кандидаты",
         "Доля покрытия", "Пояснение"],
        _candidate_gate_rows(inference), percent_columns={4},
    )
    _write_sheet(
        workbook, "По классам источника",
        ["Класс источника", "Примеров в проверке", "Точность прогнозов",
         "Полнота", "F1"],
        _source_class_rows(inference), percent_columns={3, 4, 5},
    )
    confusion_headers, confusion_rows = _source_confusion(inference)
    _write_sheet(workbook, "Матрица ошибок", confusion_headers, confusion_rows)
    _write_sheet(workbook, "Пояснения", ["Тема", "Пояснение"], _notes_rows(inference))

    workbook.properties.title = "FinOlimp — отчёт Stage 2"
    workbook.properties.subject = "Временная проверка моделей Stage 2"
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output.name}.", suffix=".tmp", dir=output.parent,
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        workbook.save(temporary)
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)
    return output


def export_stage2_workbook(stage2_dir: Path, run_id: str,
                           output: Path | None = None) -> Path:
    """Load one run's JSON reports and write its Russian-language workbook."""
    stage2_dir = Path(stage2_dir)
    inference_path = stage2_dir / "stage2_inference_report.json"
    training_path = stage2_dir / "train" / "training_report.json"
    inference = json.loads(inference_path.read_text(encoding="utf-8"))
    training = json.loads(training_path.read_text(encoding="utf-8"))
    return build_stage2_workbook(
        inference, training, run_id, output or stage2_dir / REPORT_NAME,
    )
