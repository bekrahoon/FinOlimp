"""Create a presentation-only workbook from one run's allowlisted CSV reports."""

from __future__ import annotations

import csv
import os
import re
import tempfile
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .report_dictionary import (
    describe_column,
    quality_check_meaning,
    quality_check_title,
    report_column_title,
    report_sheet_title,
    review_issue_meaning,
    review_issue_title,
)

MAX_EXCEL_ROWS = 1_048_576
MAX_EXCEL_COLUMNS = 16_384
MAX_COLUMN_WIDTH = 78
MIN_COLUMN_WIDTH = 14
_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_INVALID_XML_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_TEXT_COLUMNS = {
    "agent", "employee", "agents", "date", "identifier", "comment", "source",
    "issue", "evidence", "source_rows", "check", "key", "sample", "file",
    "raw", "norm", "label", "prefix", "top_factors", "alias_of", "month",
    "creator", "kind",
}

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(name="Aptos", bold=True, color="FFFFFF")
_HEADER_ALIGNMENT = Alignment(vertical="center", wrap_text=True)
_BODY_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
_REPORT_EXTRA_COLUMNS = {
    "quality_report.csv": ("Пояснение проверки",),
    "review_candidates.csv": ("Пояснение причины",),
}


def _sheet_title(filename: str, used: set[str]) -> str:
    base = _INVALID_SHEET_CHARS.sub("_", Path(filename).stem).strip("' ") or "Report"
    candidate = base[:31]
    suffix = 2
    while candidate.casefold() in used:
        ending = f"_{suffix}"
        candidate = f"{base[:31 - len(ending)]}{ending}"
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def _excel_text(value: str) -> str:
    # Keep report values literal (including values beginning with '=') and replace
    # only characters forbidden by the XLSX/XML format. This exporter never edits CSV.
    return _INVALID_XML_CHARS.sub("\uFFFD", value)


def _cell_value(header: str, raw: str) -> str | int | float | bool | None:
    value = _excel_text(raw)
    normalized_header = header.strip().casefold()
    if value == "":
        return None
    if normalized_header in _TEXT_COLUMNS or "identifier" in normalized_header \
            or normalized_header.endswith("_id") or value.strip() != value:
        return value
    if value.casefold() in {"true", "false"}:
        return value.casefold() == "true"
    try:
        number = Decimal(value)
    except InvalidOperation:
        return value
    if not number.is_finite():
        return value
    digits = "".join(str(digit) for digit in number.as_tuple().digits)
    significant_digits = len(digits.lstrip("0").rstrip("0"))
    if significant_digits > 15 or number.adjusted() >= 15:
        return value
    if number == number.to_integral_value():
        return int(number)
    return float(number)


def _validated_report_paths(
    reports_dir: Path,
    report_files: Iterable[str],
    output: Path,
) -> tuple[Path, list[tuple[str, Path]]]:
    reports_dir = reports_dir.resolve(strict=True)
    output_parent = output.parent.resolve(strict=True)
    if output_parent != reports_dir:
        raise ValueError("Книга должна сохраняться рядом с отчётами одного запуска.")
    if output.is_symlink():
        raise ValueError("Файл Excel-отчёта не может быть символической ссылкой.")

    paths: list[tuple[str, Path]] = []
    for filename in report_files:
        if Path(filename).name != filename or "\\" in filename or not filename.endswith(".csv"):
            raise ValueError("Разрешены только имена CSV-отчётов без путей.")
        path = reports_dir / filename
        if path.is_symlink() or not path.is_file() or path.resolve().parent != reports_dir:
            raise ValueError(f"В запуске отсутствует безопасный CSV-отчёт: {filename}.")
        paths.append((filename, path))
    if not paths:
        raise ValueError("Нет CSV-отчётов для создания книги.")
    return reports_dir, paths


def _save_atomically(workbook: Workbook, reports_dir: Path, output: Path) -> Path:
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output.name}.", suffix=".tmp", dir=reports_dir,
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        workbook.save(temporary)
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)
    return output


def _report_display_row(
    filename: str,
    headers: list[str],
    raw_row: list[str],
    *,
    is_header: bool,
) -> list[str | int | float | bool | None]:
    if is_header:
        return [report_column_title(filename, header) for header in headers] \
            + list(_REPORT_EXTRA_COLUMNS.get(filename, ()))

    if len(raw_row) != len(headers):
        raise ValueError(f"Повреждена структура CSV-отчёта {filename}.")

    values: list[str | int | float | bool | None] = []
    for header, raw_value in zip(headers, raw_row):
        if filename == "quality_report.csv" and header == "check":
            value = quality_check_title(raw_value)
        elif filename == "review_candidates.csv" and header == "issue":
            value = review_issue_title(raw_value)
        else:
            value = _cell_value(header, raw_value)
        values.append(value)

    if filename == "quality_report.csv":
        values.append(quality_check_meaning(raw_row[headers.index("check")]))
    elif filename == "review_candidates.csv":
        values.append(review_issue_meaning(raw_row[headers.index("issue")]))
    return values


def _measure_report_widths(filename: str, path: Path) -> tuple[list[str], list[int], int]:
    """Measure display values before streaming rows so Excel keeps column widths."""
    headers: list[str] = []
    widths: list[int] = []
    row_count = 0
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        for row_count, raw_row in enumerate(csv.reader(source), start=1):
            if row_count > MAX_EXCEL_ROWS:
                raise ValueError(f"Отчёт {filename} превышает лимит Excel.")
            if row_count == 1:
                headers = raw_row
            display_row = _report_display_row(
                filename, headers, raw_row, is_header=(row_count == 1),
            )
            if len(display_row) > MAX_EXCEL_COLUMNS:
                raise ValueError(f"Отчёт {filename} превышает лимит Excel.")
            if len(widths) < len(display_row):
                widths.extend([0] * (len(display_row) - len(widths)))
            for index, value in enumerate(display_row):
                text = _excel_text("") if value is None else _excel_text(str(value))
                widths[index] = max(
                    widths[index], min(len(text), MAX_COLUMN_WIDTH - 2),
                )
    return headers, widths, row_count


def export_reports_workbook(
    reports_dir: Path,
    report_files: Iterable[str],
    output: Path,
) -> Path:
    """Export each named CSV as its own Excel sheet without altering the CSVs."""
    reports_dir, paths = _validated_report_paths(reports_dir, report_files, output)

    workbook = Workbook(write_only=True)
    used_titles: set[str] = set()
    for filename, path in paths:
        headers, column_widths, measured_rows = _measure_report_widths(filename, path)
        worksheet = workbook.create_sheet(
            _sheet_title(report_sheet_title(filename), used_titles),
        )
        worksheet.freeze_panes = "A2"
        worksheet.row_dimensions[1].height = 30
        for index, width in enumerate(column_widths, start=1):
            worksheet.column_dimensions[get_column_letter(index)].width = max(
                MIN_COLUMN_WIDTH, min(width + 2, MAX_COLUMN_WIDTH),
            )
        row_count = 0
        column_count = 0
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            for row_count, raw_row in enumerate(csv.reader(source), start=1):
                row = _report_display_row(
                    filename, headers, raw_row, is_header=(row_count == 1),
                )
                if row_count > MAX_EXCEL_ROWS or len(row) > MAX_EXCEL_COLUMNS:
                    raise ValueError(f"Отчёт {filename} превышает лимит Excel.")
                column_count = max(column_count, len(row))
                cells = []
                for column_index, value in enumerate(row):
                    if isinstance(value, str):
                        value = _excel_text(value)
                    cell = WriteOnlyCell(worksheet, value=value)
                    # Never let untrusted report text become an Excel formula.
                    if isinstance(value, str):
                        cell.data_type = "s"
                        cell.number_format = "@"
                    if row_count == 1:
                        cell.fill = _HEADER_FILL
                        cell.font = _HEADER_FONT
                        cell.alignment = _HEADER_ALIGNMENT
                    else:
                        cell.alignment = _BODY_ALIGNMENT
                    cells.append(cell)
                    display_value = "" if value is None else str(value)
                    if column_index == len(column_widths):
                        column_widths.append(0)
                    column_widths[column_index] = max(
                        column_widths[column_index],
                        min(len(display_value), MAX_COLUMN_WIDTH - 2),
                    )
                worksheet.append(cells)

        if row_count:
            if column_count:
                worksheet.auto_filter.ref = (
                    f"A1:{get_column_letter(column_count)}{row_count}"
                )
        if row_count != measured_rows:
            raise ValueError(f"Отчёт {filename} изменился во время создания книги.")

    return _save_atomically(workbook, reports_dir, output)


def export_report_readme(
    reports_dir: Path,
    report_files: Iterable[str],
    output: Path,
) -> Path:
    """Write one data-dictionary sheet per CSV, with one row per CSV column."""
    reports_dir, paths = _validated_report_paths(reports_dir, report_files, output)
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_titles: set[str] = set()

    for filename, path in paths:
        worksheet = workbook.create_sheet(
            _sheet_title(report_sheet_title(filename), used_titles),
        )
        worksheet.freeze_panes = "A2"
        worksheet.row_dimensions[1].height = 30
        worksheet.column_dimensions["A"].width = 30
        worksheet.column_dimensions["B"].width = 36
        worksheet.column_dimensions["C"].width = 92
        headers = []
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            headers = next(csv.reader(source), [])
        if not headers:
            raise ValueError(f"В отчёте {filename} отсутствует строка заголовков.")

        worksheet.append([
            "Имя поля отчёта", "Русское название поля", "Пояснение для бухгалтера",
        ])
        for cell in worksheet[1]:
            cell.data_type = "s"
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = _HEADER_ALIGNMENT

        for column in headers:
            description = describe_column(filename, column)
            if not description:
                raise ValueError(
                    f"Для колонки {column!r} в отчёте {filename} ещё нет описания."
                )
            column_title = report_column_title(filename, column)
            worksheet.append([column, column_title, description])
            for cell in worksheet[worksheet.max_row]:
                cell.data_type = "s"
                cell.number_format = "@"
                cell.alignment = _BODY_ALIGNMENT

        worksheet.auto_filter.ref = f"A1:C{len(headers) + 1}"

    return _save_atomically(workbook, reports_dir, output)
