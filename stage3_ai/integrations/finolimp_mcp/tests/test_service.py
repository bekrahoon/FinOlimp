import asyncio
import csv
import json
import os
import re
import stat
from pathlib import Path

import pytest
from openpyxl import load_workbook

from finolimp_mcp import service
from finolimp_mcp import settings
from finolimp_mcp.excel_reports import export_report_readme, export_reports_workbook
from finolimp_mcp.report_dictionary import (
    REPORT_COLUMN_DESCRIPTIONS,
    REPORT_COLUMN_TITLES,
    REPORT_SHEET_TITLES,
    QUALITY_CHECK_MEANINGS,
    QUALITY_CHECK_TITLES,
    REVIEW_ISSUE_MEANINGS,
    REVIEW_ISSUE_TITLES,
)


def test_paths_resolve_stage1_at_workspace_root_and_keep_existing_runtime():
    workspace_root = settings.STAGE3_ROOT.parent

    assert settings.PROJECT_ROOT == workspace_root
    assert settings.STAGE1_SRC.is_dir()
    assert settings.CURRENT_INPUT.is_dir()
    assert settings.RUNTIME_ROOT == settings.STAGE3_ROOT / ".finolimp" / "runtime"
    assert (settings.RUNTIME_ROOT / "runs" / "run-d884555ac1874147986fc3e181e79b9a").is_dir()


def test_export_reports_workbook_localizes_columns_and_quality_codes(tmp_path):
    reports = tmp_path / "reports"
    reports.mkdir()
    (reports / "discrepancies.csv").write_text(
        '\ufeffagent,date,identifier,sum_1c,comment,source\nAgent,2026-01-01,001,2,=1+1,выяснить\n',
        encoding="utf-8",
    )
    (reports / "quality_report.csv").write_text(
        "check,count,sample\nbalance_chain_break,3,agent-a\n", encoding="utf-8",
    )
    (reports / "review_candidates.csv").write_text(
        "issue,agent,date,identifier,amount_delta,evidence,source_rows\n"
        "possible_fx_difference,Agent,2026-01-01,T-1,2,проверить курс,acts:1\n",
        encoding="utf-8",
    )
    output = reports / settings.REPORT_WORKBOOK

    result = export_reports_workbook(
        reports, ("discrepancies.csv", "quality_report.csv", "review_candidates.csv"), output,
    )

    assert result == output
    workbook = load_workbook(output, read_only=False, data_only=False)
    try:
        assert workbook.sheetnames == [
            "Расхождения", "Качество данных", "Кандидаты на проверку",
        ]
        risk_sheet = workbook["Расхождения"]
        assert [risk_sheet.cell(1, col).value for col in range(1, 7)] == [
            "Субагент", "Дата", "Билеты и документы", "Сумма по 1С",
            "Комментарий сверки", "Предварительная причина",
        ]
        assert risk_sheet["A2"].value == "Agent"
        assert risk_sheet["D2"].value == 2
        assert risk_sheet["D2"].data_type == "n"
        assert risk_sheet["C2"].value == "001"
        assert risk_sheet["C2"].data_type == "s"
        assert risk_sheet["E2"].value == "=1+1"
        assert risk_sheet["E2"].data_type == "s"
        assert risk_sheet["E2"].alignment.wrap_text is True
        assert risk_sheet.freeze_panes == "A2"
        assert risk_sheet.auto_filter.ref == "A1:F2"
        quality_sheet = workbook["Качество данных"]
        assert [quality_sheet.cell(1, col).value for col in range(1, 5)] == [
            "Проверка", "Количество случаев", "Пример и детали", "Пояснение проверки",
        ]
        assert quality_sheet["A2"].value == "Нарушена последовательность баланса ETM"
        assert quality_sheet["B2"].value == 3
        assert "не согласуется с предыдущим остатком" in quality_sheet["D2"].value
        assert quality_sheet["D2"].alignment.wrap_text is True
        assert quality_sheet.column_dimensions["D"].width > 52
        candidates = workbook["Кандидаты на проверку"]
        assert candidates["A1"].value == "Причина проверки"
        assert candidates["A2"].value == "Возможное расхождение из-за валютного курса"
        assert "валютного курса" in candidates["H2"].value
    finally:
        workbook.close()


def test_export_reports_workbook_rejects_missing_files_and_paths(tmp_path):
    reports = tmp_path / "reports"
    reports.mkdir()
    output = reports / settings.REPORT_WORKBOOK

    with pytest.raises(ValueError, match="без путей"):
        export_reports_workbook(reports, ("../outside.csv",), output)
    with pytest.raises(ValueError, match="отсутствует"):
        export_reports_workbook(reports, ("missing.csv",), output)


def test_localize_report_code_values_preserves_headers_and_mcp_codes(tmp_path):
    reports = tmp_path / "reports"
    reports.mkdir()
    quality = reports / "quality_report.csv"
    quality.write_text(
        "check,key,count,sample\nbalance_chain_break,Agent A,1,row 12\n",
        encoding="utf-8",
    )
    candidates = reports / "review_candidates.csv"
    candidates.write_text(
        "issue,agent,date,identifier,amount_delta,evidence,source_rows\n"
        "possible_fx_difference,Agent A,2026-01-01,T-1,2,rate,acts:12\n",
        encoding="utf-8",
    )

    service._localize_report_code_values(reports)

    with quality.open(encoding="utf-8-sig", newline="") as source:
        quality_reader = csv.DictReader(source)
        quality_headers = quality_reader.fieldnames
        quality_rows = list(quality_reader)
    assert quality_headers == ["check", "key", "count", "sample"]
    assert quality_rows[0]["check"] == QUALITY_CHECK_MEANINGS["balance_chain_break"]
    assert service._read_rows(quality)[0]["check"] == "balance_chain_break"

    with candidates.open(encoding="utf-8-sig", newline="") as source:
        candidate_reader = csv.DictReader(source)
        candidate_headers = candidate_reader.fieldnames
        candidate_rows = list(candidate_reader)
    assert candidate_headers == [
        "issue", "agent", "date", "identifier", "amount_delta", "evidence", "source_rows",
    ]
    assert candidate_rows[0]["issue"] == REVIEW_ISSUE_MEANINGS["possible_fx_difference"]
    assert candidate_rows[0]["agent"] == "Agent A"
    assert service._read_rows(candidates)[0]["issue"] == "possible_fx_difference"


def test_export_report_readme_creates_one_description_sheet_per_csv(tmp_path):
    reports = tmp_path / "reports"
    reports.mkdir()
    (reports / "agent_risk.csv").write_text(
        "agent,rows,err_rate\n", encoding="utf-8",
    )
    (reports / "quality_report.csv").write_text(
        "check,count,sample\n", encoding="utf-8",
    )
    (reports / "review_candidates.csv").write_text(
        "issue,agent,date,identifier,amount_delta,evidence,source_rows\n",
        encoding="utf-8",
    )
    output = reports / settings.REPORT_README

    result = export_report_readme(
        reports, ("agent_risk.csv", "quality_report.csv", "review_candidates.csv"), output,
    )

    assert result == output
    workbook = load_workbook(output, read_only=False, data_only=False)
    try:
        assert workbook.sheetnames == [
            "Риск субагентов", "Качество данных", "Кандидаты на проверку",
        ]
        risk_sheet = workbook["Риск субагентов"]
        assert [risk_sheet.cell(1, col).value for col in range(1, 4)] == [
            "Имя поля отчёта", "Русское название поля", "Пояснение для бухгалтера",
        ]
        assert [risk_sheet.cell(2, col).value for col in range(1, 4)] == [
            "agent", "Субагент", "Каноническое имя субагента.",
        ]
        assert "Количество строк" in risk_sheet["C3"].value
        quality_sheet = workbook["Качество данных"]
        assert quality_sheet["A4"].value == "sample"
        assert quality_sheet["B4"].value == "Пример и детали"
        assert quality_sheet["C4"].value
        assert workbook["Кандидаты на проверку"]["B2"].value == "Причина проверки"
    finally:
        workbook.close()


def test_export_report_readme_requires_description_for_every_column(tmp_path):
    reports = tmp_path / "reports"
    reports.mkdir()
    (reports / "agent_risk.csv").write_text(
        "agent,new_unmapped_field\n", encoding="utf-8",
    )

    with pytest.raises(ValueError, match="ещё нет описания"):
        export_report_readme(
            reports, ("agent_risk.csv",), reports / settings.REPORT_README,
        )


def test_column_descriptions_match_checked_in_stage1_report_schemas():
    reports = settings.PROJECT_ROOT / "stage1_rules" / "reports"
    if not all((reports / name).is_file() for name in service.REPORT_FILES):
        pytest.skip("В этой копии проекта нет полного набора Stage 1 CSV-схем.")

    for filename in service.REPORT_FILES:
        with (reports / filename).open(encoding="utf-8-sig", newline="") as source:
            columns = next(csv.reader(source), [])
        assert set(columns) == set(REPORT_COLUMN_DESCRIPTIONS[filename]), filename
        assert set(columns) == set(REPORT_COLUMN_TITLES[filename]), filename

    assert set(service.REPORT_FILES) == set(REPORT_SHEET_TITLES)
    quality_report = reports / "quality_report.csv"
    with quality_report.open(encoding="utf-8-sig", newline="") as source:
        actual_checks = {row["check"] for row in csv.DictReader(source)}
    assert actual_checks <= set(QUALITY_CHECK_MEANINGS)
    assert actual_checks <= set(QUALITY_CHECK_TITLES)
    issue_report = reports / "review_candidates.csv"
    with issue_report.open(encoding="utf-8-sig", newline="") as source:
        actual_issues = {row["issue"] for row in csv.DictReader(source)}
    assert actual_issues <= set(REVIEW_ISSUE_MEANINGS)
    assert actual_issues <= set(REVIEW_ISSUE_TITLES)
    for filename, descriptions in REPORT_COLUMN_DESCRIPTIONS.items():
        for column, description in descriptions.items():
            latin_words = set(re.findall(r"\b[A-Za-z][A-Za-z0-9_]*\b", description))
            assert latin_words <= {"ETM"}, f"{filename}:{column}: {latin_words}"
    assert all(re.search(r"[А-Яа-яЁё]", title) for title in REPORT_SHEET_TITLES.values())


def test_run_reconciliation_adds_localized_workbooks_and_preserves_csv_headers(
    tmp_path, monkeypatch,
):
    dataset_id = "ds-" + "a" * 32
    dataset_dir = tmp_path / "runtime" / "datasets" / dataset_id
    dataset_dir.mkdir(parents=True)
    (dataset_dir / "manifest.json").write_text(
        json.dumps({"dataset_id": dataset_id, "files": {}}), encoding="utf-8",
    )
    monkeypatch.setattr(service, "validate_dataset", lambda *_args: {"ready": True})
    monkeypatch.setattr(service, "_freeze_tree", lambda *_args: None)

    async def fake_create_subprocess_exec(*command, **_kwargs):
        reports_dir = Path(command[command.index("--out") + 1])
        headers = {
            "discrepancies.csv": "agent,date,identifier,sum_1c,sum_etm,diff,comment,source",
            "reconciliation_summary.csv": "agent,residual",
            "quality_report.csv": "check,key,count,sample",
            "review_candidates.csv": "issue,agent,date,identifier,amount_delta,evidence,source_rows",
            "agent_actions.csv": "agent,n_rows,n_bot,n_1c,n_agent,n_vyjasnit,move_etm_by,fix_1c_by,clarify_money,clarify_n,operation_diff,residual,action_money",
            "agent_risk.csv": "agent,rows,exp_errors,exp_money,old_money_30d,fresh_rows,sum_abs_diff,err_rate,top_factors",
            "employee_risk.csv": "employee,total_rows,suspect_rows,witness_rows,agents,suspect_money,suspect_rate",
            "features.csv": "agent,date,identifier,diff",
            "label_map.csv": "file,raw,norm,label",
        }
        for filename in service.REPORT_FILES:
            (reports_dir / filename).write_text(headers[filename] + "\n", encoding="utf-8")

        class CompletedProcess:
            returncode = 0

            async def communicate(self):
                return b"mock report generation", None

        return CompletedProcess()

    monkeypatch.setattr(service.asyncio, "create_subprocess_exec", fake_create_subprocess_exec)

    result = asyncio.run(service.run_reconciliation(dataset_id, tmp_path / "runtime"))

    reports_dir = (
        tmp_path / "runtime" / "runs" / result["run_id"] / "reports"
    )
    assert result["status"] == "completed", result
    assert result["report_workbook"] == settings.REPORT_WORKBOOK
    assert result["available_reports"] == [
        *service.REPORT_FILES, settings.REPORT_WORKBOOK, settings.REPORT_README,
    ]
    workbook = load_workbook(reports_dir / settings.REPORT_WORKBOOK, read_only=True)
    try:
        assert workbook.sheetnames == [REPORT_SHEET_TITLES[filename] for filename in service.REPORT_FILES]
    finally:
        workbook.close()
    readme = load_workbook(reports_dir / settings.REPORT_README, read_only=True)
    try:
        assert readme.sheetnames == [REPORT_SHEET_TITLES[filename] for filename in service.REPORT_FILES]
        assert readme["Признаки расхождений"]["A2"].value == "agent"
        assert readme["Признаки расхождений"]["B2"].value
    finally:
        readme.close()
    summary = service.get_run_summary(result["run_id"], tmp_path / "runtime")
    assert summary["report_workbook"] == settings.REPORT_WORKBOOK
    assert summary["report_readme"] == settings.REPORT_README
    assert settings.REPORT_WORKBOOK in summary["available_reports"]


def _write_inputs(folder):
    folder.mkdir(parents=True)
    for filename in service.INPUT_FILES:
        (folder / filename).write_text(f"placeholder for {filename}\n", encoding="utf-8")
    (folder / "reconcile.yaml").write_text("tolerance: 0.01\n", encoding="utf-8")
    (folder / "unrelated.txt").write_text("not part of the dataset", encoding="utf-8")


def _write_report(path, rows, fieldnames):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def _completed_run(runtime_root):
    run_id = "run-" + "a" * 32
    run_dir = runtime_root / "runs" / run_id
    reports = run_dir / "reports"
    reports.mkdir(parents=True)
    (run_dir / "manifest.json").write_text(json.dumps({
        "run_id": run_id,
        "dataset_id": "ds-" + "b" * 32,
        "status": "completed",
        "pipeline_fingerprint": "f" * 64,
        "started_at": "2026-01-01T00:00:00+00:00",
        "finished_at": "2026-01-01T00:01:00+00:00",
        "reports": list(service.REPORT_FILES),
    }), encoding="utf-8")
    _write_report(reports / "discrepancies.csv", [
        {"agent": "Agent A", "date": "2026-01-01", "identifier": "T-1", "sum_1c": "10",
         "sum_etm": "9", "diff": "1", "comment": "Проверить первичку", "source": "выяснить"},
        {"agent": "Agent B", "date": "2026-01-02", "identifier": "T-2", "sum_1c": "20",
         "sum_etm": "20", "diff": "0", "comment": "No diff", "source": "1С"},
    ], ["agent", "date", "identifier", "sum_1c", "sum_etm", "diff", "comment", "source"])
    _write_report(reports / "reconciliation_summary.csv", [
        {"agent": "Agent A", "residual": "0.02"}, {"agent": "Agent B", "residual": "0"},
    ], ["agent", "residual"])
    _write_report(reports / "agent_actions.csv", [
        {"agent": "Agent A", "move_etm_by": "5.25", "fix_1c_by": "-2", "clarify_money": "7.5", "clarify_n": "2"},
        {"agent": "Agent B", "move_etm_by": "0", "fix_1c_by": "0", "clarify_money": "0", "clarify_n": "0"},
    ], ["agent", "move_etm_by", "fix_1c_by", "clarify_money", "clarify_n"])
    _write_report(reports / "review_candidates.csv", [
        {"issue": "possible_fx_difference", "agent": "Agent A", "date": "2026-01-01",
         "identifier": "T-1", "amount_delta": "1", "evidence": "candidate", "source_rows": ""},
    ], ["issue", "agent", "date", "identifier", "amount_delta", "evidence", "source_rows"])
    _write_report(reports / "quality_report.csv", [
        {"check": "balance_chain_break", "key": "Agent A", "count": "2", "sample": ""},
    ], ["check", "key", "count", "sample"])
    return run_id


def test_register_current_dataset_snapshots_allowlisted_files(tmp_path):
    source = tmp_path / "input"
    runtime = tmp_path / "runtime"
    _write_inputs(source)

    result = service.register_current_dataset(source, runtime)

    snapshot = runtime / "datasets" / result["dataset_id"]
    assert result["dataset_id"].startswith("ds-")
    assert set(result["files"]) == {*service.INPUT_FILES, "reconcile.yaml"}
    assert not (snapshot / "unrelated.txt").exists()
    assert (snapshot / "manifest.json").is_file()
    assert result["files"]["acts.csv"]["sha256"]
    if os.name == "posix":
        assert stat.S_IMODE(snapshot.stat().st_mode) == 0o500
        assert stat.S_IMODE((snapshot / "acts.csv").stat().st_mode) == 0o400
        assert stat.S_IMODE((snapshot / "manifest.json").stat().st_mode) == 0o400


def test_register_current_dataset_requires_all_inputs(tmp_path):
    source = tmp_path / "input"
    source.mkdir()
    (source / "acts.csv").write_text("data", encoding="utf-8")

    with pytest.raises(ValueError, match="обязательный файл"):
        service.register_current_dataset(source, tmp_path / "runtime")


def test_validate_dataset_detects_snapshot_tampering_before_parsing(tmp_path):
    source = tmp_path / "input"
    _write_inputs(source)
    result = service.register_current_dataset(source, tmp_path / "runtime")
    snapshot_file = tmp_path / "runtime" / "datasets" / result["dataset_id"] / "acts.csv"
    snapshot_file.chmod(0o600)
    snapshot_file.write_text("changed after registration", encoding="utf-8")

    validation = service.validate_dataset(result["dataset_id"], tmp_path / "runtime")

    assert validation["ready"] is False
    assert "Создайте новый snapshot" in validation["problems"][0]


@pytest.mark.parametrize("dataset_id", ["../outside", "ds-not-an-id", "ds-" + "A" * 32])
def test_dataset_lookup_rejects_invalid_ids(tmp_path, dataset_id):
    with pytest.raises(ValueError, match="идентификатор"):
        service._dataset_dir(dataset_id, tmp_path)


def test_list_discrepancies_returns_paginated_rows_and_filters(tmp_path):
    run_id = _completed_run(tmp_path)

    page = service.list_discrepancies(run_id, source="выяснить", limit=1, runtime_root=tmp_path)
    assert page["total"] == 1
    assert page["discrepancies"][0]["discrepancy_id"] == "dsp-000001"
    assert page["discrepancies"][0]["agent"] == "Agent A"

    detail = service.get_discrepancy_detail(run_id, "dsp-000001", tmp_path)
    assert detail["diff"] == "1"
    assert detail["source_rows_available"] is False
    assert "Stage 1" in detail["limitation"]


def test_summary_keeps_manual_candidates_separate(tmp_path):
    run_id = _completed_run(tmp_path)

    summary = service.get_run_summary(run_id, tmp_path)

    assert summary["discrepancy_count"] == 2
    assert summary["discrepancies_by_source"] == {"выяснить": 1, "1С": 1}
    assert summary["agents_with_nonzero_residual"] == [{"agent": "Agent A", "residual": "0.02"}]
    assert summary["proposed_actions"]["move_etm_by"] == 5.25
    assert summary["proposed_actions"]["clarify_rows"] == 2.0
    assert "не автоматически" in summary["proposed_actions"]["notice"]
    assert summary["manual_review_candidate_count"] == 1
    assert summary["manual_review_candidates_by_type"] == {"possible_fx_difference": 1}
    assert summary["manual_review_candidate_types"][0]["reason"] == (
        "Возможное расхождение из-за валютного курса"
    )
    assert summary["quality_observations"][0]["meaning"]
    assert summary["quality_observations"][0]["check"] == (
        "Нарушена последовательность баланса ETM"
    )


def test_list_runs_sorts_by_timestamp_not_random_id(tmp_path):
    older_id = "run-" + "f" * 32
    newer_id = "run-" + "0" * 32
    runs_root = tmp_path / "runs"
    for run_id, started_at in ((older_id, "2026-01-01T00:00:00Z"),
                               (newer_id, "2026-02-01T00:00:00Z")):
        path = runs_root / run_id
        path.mkdir(parents=True)
        (path / "manifest.json").write_text(json.dumps({
            "run_id": run_id,
            "dataset_id": "ds-" + "b" * 32,
            "status": "completed",
            "started_at": started_at,
            "finished_at": started_at,
        }), encoding="utf-8")

    result = service.list_runs(runtime_root=tmp_path)

    assert [run["run_id"] for run in result["runs"]] == [newer_id, older_id]


def test_list_quality_findings_explains_observation_without_attributing_error(tmp_path):
    run_id = _completed_run(tmp_path)

    findings = service.list_quality_findings(
        run_id, check="balance_chain_break", runtime_root=tmp_path,
    )

    assert findings["total"] == 1
    assert findings["findings"][0]["count"] == "2"
    assert findings["findings"][0]["check_title"] == (
        "Нарушена последовательность баланса ETM"
    )
    assert "остаток etm" in findings["findings"][0]["meaning"].casefold()
    assert "диагностические" in findings["notice"]


def test_submit_case_for_build_persists_a_draft_without_releasing(tmp_path):
    result = service.submit_case_for_build(
        "Операция не сопоставилась с актом",
        "Нужно показать связанные строки и попросить первичку",
        runtime_root=tmp_path,
    )

    case = service.get_build_case(result["case_id"], tmp_path)
    assert result["status"] == "submitted"
    assert case["status"] == "submitted"
    assert "не изменён" in case["notice"]
    assert service.list_build_cases(runtime_root=tmp_path)["count"] == 1


def test_submit_case_for_build_rejects_invalid_link(tmp_path):
    with pytest.raises(ValueError, match="ID запуска"):
        service.submit_case_for_build("Случай", "Ожидание", discrepancy_id="dsp-000001",
                                      runtime_root=tmp_path)
